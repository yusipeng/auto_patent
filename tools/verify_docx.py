# -*- coding: utf-8 -*-
"""verify_docx.py — 交底书 docx 转写结果校验
用法: .venv_patent/Scripts/python.exe tools/verify_docx.py <out.docx>
输出: 一级章节标题(12个, 与模板对比)、图片数、公式数(OMML)、$ 残留、渲染失败标注、正文段落数、表格数

说明: 标题按 段落样式(Heading N) 与 直接 outlineLvl 双通道识别；
      模板章节含 Word 自动编号(不在文本中, 如"一、")，比对时自动剥离编号前缀。
"""
import re
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn

TEMPLATE_SECTIONS = [
    '零、术语定义和解释', '一、发明名称', '二、技术领域',
    '三、现有技术的技术方案', '四、现有技术的缺点及本申请提案要解决的技术问题',
    '五、本申请提案的技术方案的详细阐述', '六、本申请提案的关键点和欲保护点',
    '七、与第三条中最接近的现有技术相比，本申请提案有何技术优点',
    '八、发散思维以及规避方案思考',
    '本申请提案的商业价值',
    '本申请提案的侵权证据可获得性/标准进展情况',
    '其他有助于理解本申请提案的技术资料',
]

NUM_PREFIX = re.compile(r'^[零一二三四五六七八九十]+、\s*')


def heading_level(p):
    """Heading N 样式 或 直接 outlineLvl → 级别(0 起)；否则 None"""
    pPr = p._p.pPr
    if pPr is not None:
        ol = pPr.find(qn('w:outlineLvl'))
        if ol is not None:
            try:
                return int(ol.get(qn('w:val')))
            except (TypeError, ValueError):
                pass
    m = re.match(r'Heading (\d+)$', p.style.name or '')
    if m:
        return int(m.group(1)) - 1
    return None


def heading_numid(p):
    """标题段落的 Word 自动编号 numId；无则 None"""
    pPr = p._p.pPr
    if pPr is not None:
        numPr = pPr.find(qn('w:numPr'))
        if numPr is not None:
            nid = numPr.find(qn('w:numId'))
            if nid is not None:
                return nid.get(qn('w:val'))
    return None


def main():
    path = sys.argv[1]
    doc = Document(path)
    headings = []  # (text, numId)
    texts = []
    for p in doc.paragraphs:
        txt = ''.join(r.text for r in p.runs)
        if heading_level(p) == 0:
            headings.append((txt.strip(), heading_numid(p)))
        texts.append(txt)

    # 图片数: drawing/pict
    xml = doc.element.body.xml
    images = xml.count('<w:drawing>') + xml.count('<w:pict>')

    # OMML 公式数
    with zipfile.ZipFile(path) as z:
        doc_xml = z.read('word/document.xml').decode('utf-8')
    n_omath = doc_xml.count('<m:oMath>')
    n_omathpara = doc_xml.count('<m:oMathPara>')

    # $ 残留 (字面美元符号在文本节点中)
    dollar_samples = [t[:80] for t in texts if '$' in t]

    # 渲染失败标注
    fail_markers = [t[:120] for t in texts
                    if any(k in t for k in ('[公式渲染失败', '[Mermaid 渲染失败',
                                             '[图片缺失', '[图片无法嵌入', '[图渲染失败'))]

    # 一级章节标题与模板对比（剥离自动编号前缀后比对）
    hset = {NUM_PREFIX.sub('', h[0]).strip() for h in headings}
    matched = [s for s in TEMPLATE_SECTIONS if NUM_PREFIX.sub('', s).strip() in hset]
    missing = [s for s in TEMPLATE_SECTIONS if s not in matched]

    print('=== 一级章节标题 (Heading 1) ===')
    for t, nid in headings:
        tag = f'numId={nid}' if nid else '无编号'
        print(f'   [{tag:>9}] {t}')
    print(f'共 {len(headings)} 个; 与模板匹配 {len(matched)}/{len(TEMPLATE_SECTIONS)}')
    if missing:
        print('缺失:', missing)
    print()
    print(f'图片数(drawing+pict): {images}')
    print(f'OMML公式数: <m:oMath>={n_omath}  <m:oMathPara>={n_omathpara}')
    print(f'字面$残留段落: {len(dollar_samples)}')
    for s in dollar_samples:
        print('  [残留]', s)
    print(f'渲染失败标注: {len(fail_markers)}')
    for s in fail_markers:
        print('  [FAIL]', s)
    print(f'正文段落数(含标题): {len(texts)}')
    print(f'表格数: {len(doc.tables)}')


if __name__ == '__main__':
    main()
