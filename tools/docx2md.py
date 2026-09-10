# -*- coding: utf-8 -*-
"""docx2md.py — 已有交底书 docx -> 流水线 md 源（导入工具）
用法: .venv_patent/Scripts/python.exe tools/docx2md.py <in.docx> <out.md> [--keep-hints]
- 按模板结构提取: 封面字段、零~八章节 + 3 无编号节、表格、内嵌图片(存到 media_dir)
- 兼容 Word 自动编号: 标题/列表编号自 numbering.xml 合成回文字
  ("一、二、…"、子节 "3.1/5.1"、列表 "1."/"（1）"), 由引用计数器按文档顺序推进
- 公式(OMML 原生公式): 借 pandoc 逐段还原为 $...$ 行内 / ```latex 块
- 模板说明性文字(【...】占位提示)默认剥离, --keep-hints 保留
"""
import sys, os, re
import docx
from docx.oxml.ns import qn
from lxml import etree

M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
MATH_O = '{' + M_NS + '}oMath'

# 交底书章节标题 (与 md2docx.SECTION_TITLES 保持一致)
SECTION_TITLES = [
    '零、术语定义和解释', '一、发明名称', '二、技术领域',
    '三、现有技术的技术方案', '四、现有技术的缺点及本申请提案要解决的技术问题',
    '五、本申请提案的技术方案的详细阐述', '六、本申请提案的关键点和欲保护点',
    '七、与第三条中最接近的现有技术相比，本申请提案有何技术优点',
    '八、发散思维以及规避方案思考',
    '本申请提案的商业价值',
    '本申请提案的侵权证据可获得性/标准进展情况',
    '其他有助于理解本申请提案的技术资料',
]
# 归一化(去空白) 查找表: 带/不带"一、"编号均可命中 -> 规范全称
SECTION_LOOKUP = {}
for _t in SECTION_TITLES:
    SECTION_LOOKUP[re.sub(r'\s+', '', _t)] = _t
    SECTION_LOOKUP[re.sub(r'\s+', '', re.sub(r'^[一二三四五六七八九十]+、', '', _t))] = _t


def iter_block_items(doc):
    """按 body 顺序产出 ('p', paragraph_element) / ('tbl', table_element)"""
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split('}')[1]
        if tag == 'p':
            yield ('p', child)
        elif tag == 'tbl':
            yield ('tbl', child)


def para_text(p_el):
    return ''.join(t.text or '' for t in p_el.iter(qn('w:t')))


def is_hint(text):
    return bool(re.match(r'^【[^】]*】$', text.strip())) or text.strip().startswith('【')


def extract_images_from_para(doc, p_el, media_dir, counter):
    """提取段落内的内嵌图片, 返回 md 图片行列表"""
    lines = []
    blips = p_el.findall('.//' + qn('a:blip'))
    for blip in blips:
        rid = blip.get(qn('r:embed'))
        if not rid:
            continue
        try:
            part = doc.part.related_parts[rid]
        except KeyError:
            continue
        counter[0] += 1
        fname = f'img_{counter[0]:02d}.png'
        os.makedirs(media_dir, exist_ok=True)
        full = os.path.join(media_dir, fname)
        with open(full, 'wb') as f:
            f.write(part.blob)
        # 引用路径相对 md 文件所在目录(随 md 一起流转, 不依赖 CWD)
        rel = os.path.relpath(full, os.path.dirname(os.path.abspath(media_dir)) )
        lines.append(f'![原图{counter[0]}]({rel.replace(os.sep, "/")})')
    return lines


def tbl_to_md(t_el):
    rows = t_el.findall(qn('w:tr'))
    if not rows:
        return []
    def row_cells(tr):
        return [''.join(t.text or '' for t in tc.iter(qn('w:t'))).strip()
                for tc in tr.findall(qn('w:tc'))]
    md = []
    first = row_cells(rows[0])
    md.append('| ' + ' | '.join(first) + ' |')
    md.append('|' + '---|' * len(first))
    for tr in rows[1:]:
        cells = row_cells(tr)
        # 对齐列数
        while len(cells) < len(first):
            cells.append('')
        md.append('| ' + ' | '.join(cells[:len(first)]) + ' |')
    return md


# ---------------- Word 自动编号 -> 文字合成 ----------------

_ZH = '零一二三四五六七八九'


def _zh_num(n):
    if n < 10:
        return _ZH[n]
    if n == 10:
        return '十'
    if n < 20:
        return '十' + _ZH[n % 10]
    if n < 100:
        s = _ZH[n // 10] + '十'
        if n % 10:
            s += _ZH[n % 10]
        return s
    return str(n)


class NumberSynth:
    """把 Word 自动编号(标题/列表)按文档顺序合成回文字前缀。
    行内编号定义来自 numbering.xml; number_of() 必须对每个段落按文档顺序调用一次,
    即使不使用其返回值, 以保证计数器与文档一致。"""

    def __init__(self, doc):
        self.counters = {}
        self.lvl = {}   # numId -> (lvlText, numFmt, suff, start)
        try:
            el = doc.part.numbering_part.element
        except Exception:
            return
        abs_map = {}
        for a in el.findall(qn('w:abstractNum')):
            lvl = a.find(qn('w:lvl'))
            if lvl is None:
                continue
            lt = lvl.find(qn('w:lvlText'))
            nf = lvl.find(qn('w:numFmt'))
            sf = lvl.find(qn('w:suff'))
            st = lvl.find(qn('w:start'))
            abs_map[a.get(qn('w:abstractNumId'))] = (
                lt.get(qn('w:val')) if lt is not None else '',
                nf.get(qn('w:val')) if nf is not None else 'decimal',
                sf.get(qn('w:val')) if sf is not None else 'tab',
                int(st.get(qn('w:val'))) if st is not None else 1)
        for n in el.findall(qn('w:num')):
            ref = n.find(qn('w:abstractNumId'))
            if ref is None:
                continue
            self.lvl[n.get(qn('w:numId'))] = abs_map.get(
                ref.get(qn('w:val')), ('', 'decimal', 'tab', 1))

    def number_of(self, p_el):
        """返回 (编号文字, 分隔符) 或 None。编号文字可能为空串(无可见编号的列表)。"""
        ppr = p_el.find(qn('w:pPr'))
        if ppr is None:
            return None
        npr = ppr.find(qn('w:numPr'))
        if npr is None:
            return None
        nid_el = npr.find(qn('w:numId'))
        if nid_el is None:
            return None
        nid = nid_el.get(qn('w:val'))
        if nid not in self.lvl:
            return None
        lvl_text, num_fmt, suff, start = self.lvl[nid]
        c = start + self.counters.get(nid, 0)
        self.counters[nid] = self.counters.get(nid, 0) + 1
        if not lvl_text:
            return ('', '')
        if '%1' not in lvl_text:
            return (lvl_text, ' ')
        disp = lvl_text.replace('%1', _zh_num(c) if num_fmt == 'chineseCounting' else str(c))
        sep = ' ' if suff in ('space', 'tab', '') else ''
        return (disp, sep)


# ---------------- OMML 公式 -> LaTeX (pandoc 单段还原, 尽力而为) ----------------

_CT = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
       '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
       '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
       '<Default Extension="xml" ContentType="application/xml"/>'
       '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
       '</Types>')
_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
         '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
         '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
         '</Relationships>')
_DOC_HEAD = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<w:document xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
             'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
             'xmlns:o="urn:schemas-microsoft-com:office:office" '
             'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
             'xmlns:m="' + M_NS + '" '
             'xmlns:v="urn:schemas-microsoft-com:vml" '
             'xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" '
             'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
             'xmlns:w10="urn:schemas-microsoft-com:office:word" '
             'xmlns:w="' + W_NS + '" '
             'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
             'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" '
             'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
             'xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk" '
             'xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" '
             'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
             'mc:Ignorable="w14 w15 wp14"><w:body>')
_DOC_TAIL = '</w:body></w:document>'


def _unescape_outside_math(s):
    """还原 pandoc markdown 输出中数学段以外的反斜杠转义(\\" -> ", \\> -> > 等);
    $...$ 数学段保持原样(反斜杠是 LaTeX)。"""
    parts = re.split(r'(\$[^$\n]+\$)', s)
    for i, seg in enumerate(parts):
        if seg.startswith('$') and seg.endswith('$') and len(seg) > 2:
            continue
        parts[i] = re.sub(r'\\([^\w\s])', r'\1', seg)
    return ''.join(parts)


def convert_para_math(p_el, cache):
    """单段 OMML -> markdown 文本(含 $LaTeX$), 失败返回 None"""
    import zipfile, tempfile as _tf, uuid, shutil, subprocess
    frag = etree.tostring(p_el, encoding='unicode')
    if frag in cache:
        return cache[frag]
    out = None
    wd = os.path.join(_tf.gettempdir(), 'd2m_' + uuid.uuid4().hex)
    try:
        os.makedirs(os.path.join(wd, 'word'), exist_ok=True)
        os.makedirs(os.path.join(wd, '_rels'), exist_ok=True)
        el = etree.fromstring(frag)
        doc_xml = _DOC_HEAD + etree.tostring(el, encoding='unicode') + _DOC_TAIL
        with open(os.path.join(wd, 'word', 'document.xml'), 'w', encoding='utf-8') as f:
            f.write(doc_xml)
        with open(os.path.join(wd, '_rels', '.rels'), 'w', encoding='utf-8') as f:
            f.write(_RELS)
        with open(os.path.join(wd, '[Content_Types].xml'), 'w', encoding='utf-8') as f:
            f.write(_CT)
        dx = os.path.join(wd, 's.docx')
        with zipfile.ZipFile(dx, 'w', zipfile.ZIP_DEFLATED) as zf:
            zf.write(os.path.join(wd, 'word', 'document.xml'), 'word/document.xml')
            zf.write(os.path.join(wd, '_rels', '.rels'), '_rels/.rels')
            zf.write(os.path.join(wd, '[Content_Types].xml'), '[Content_Types].xml')
        r = subprocess.run(['pandoc', dx, '-t', 'markdown', '--wrap=none'],
                           capture_output=True, text=True, timeout=60)
        if r.returncode == 0 and r.stdout.strip():
            out = r.stdout.strip()
    except Exception:
        out = None
    finally:
        shutil.rmtree(wd, ignore_errors=True)
    cache[frag] = out
    return out


# ---------------- 主转换 ----------------

def convert(docx_path, out_md, keep_hints=False):
    doc = docx.Document(docx_path)
    synth = NumberSynth(doc)
    style_names = {}
    try:
        for st in doc.styles:
            style_names[st.style_id] = st.name
    except Exception:
        pass

    def heading_level(p_el):
        ppr = p_el.find(qn('w:pPr'))
        if ppr is not None:
            ps = ppr.find(qn('w:pStyle'))
            if ps is not None:
                nm = (style_names.get(ps.get(qn('w:val'))) or '').strip()
                m = re.match(r'(?i)^heading\s*(\d)', nm)
                if m:
                    return int(m.group(1))
            ol = ppr.find(qn('w:outlineLvl'))
            if ol is not None:
                try:
                    return int(ol.get(qn('w:val'))) + 1
                except Exception:
                    pass
        return None

    media_dir = os.path.join(os.path.dirname(out_md) or '.', 'media')
    counter = [0]
    math_cache = {}
    out = []
    # 封面发明名称: 记录备用, 写入正文「一、发明名称」标题处 (md2docx 回填封面时从那里读取)
    cover_name = None
    if doc.tables:
        cover = doc.tables[0]
        for row in cover.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0] == '发明名称':
                cover_name = cells[1]
                break
    name_emitted = False
    seen_headings = {'一、发明名称'}  # 封面已写入, 正文重复标题跳过
    in_hint = False  # 跨段 【...】 提示区
    tbl_idx = -1
    items = list(iter_block_items(doc))

    def has_content_after(idx):
        """判断某标题后是否存在正文内容段(至下一章节标题为止)"""
        for k2, e2 in items[idx + 1:]:
            if k2 != 'p':
                continue
            t2 = para_text(e2).strip()
            if not t2:
                continue
            if re.sub(r'\s+', '', t2) in SECTION_LOOKUP:
                return False
            return True
        return False

    for idx, (kind, el) in enumerate(items):
        if kind == 'p':
            text = para_text(el).strip()
            num = synth.number_of(el)   # 必须先调用(推进计数器), 与文档顺序一致
            img_lines = extract_images_from_para(doc, el, media_dir, counter)
            if img_lines:
                out.append('')
                out.extend(img_lines)
                out.append('')
            # OMML 公式段: pandoc 还原为 markdown(含 $LaTeX$)
            if el.find('.//' + MATH_O) is not None:
                raw = convert_para_math(el, math_cache)
                if raw:
                    s = raw.strip()
                    if len(s) > 4 and s.startswith('$$') and s.endswith('$$') and s.count('$$') == 2:
                        # 整段为独立公式 -> latex 代码块
                        out.append('')
                        out.append('```latex')
                        out.append(s[2:-2].strip())
                        out.append('```')
                        out.append('')
                        continue
                    text = _unescape_outside_math(s).strip()
            if not text:
                continue
            # 封面杂项行跳过
            if text in ('公司专利申请', '技术交底书', '〔集团名称〕', '〔集团名称〕', '附件3-发明&实用新型交底书模板'):
                continue
            if text.startswith('附件3-'):
                continue
            # 章节标题(兼容 Word 自动编号: 标题文字可能不含"一、"; 以及旧版文字编号)
            canon = SECTION_LOOKUP.get(re.sub(r'\s+', '', text)) if len(text) < 80 else None
            if canon:
                if canon in seen_headings:
                    # 已见过的标题(封面区已含名称): 视情况补入封面名称, 避免与正文内容重复
                    if canon == '一、发明名称' and not name_emitted:
                        out.append('')
                        out.append('# 一、发明名称')
                        out.append('')
                        if cover_name and not has_content_after(idx):
                            out.append(cover_name)
                            out.append('')
                        name_emitted = True
                    continue
                seen_headings.add(canon)
                out.append('')
                out.append(f'# {canon}')
                out.append('')
                continue
            # 子节标题(## / ### / ####), 自动编号合成回文字
            lvl = heading_level(el)
            if lvl == 1 and len(text) < 80:
                # 章节表之外的 H1 标题: 原样保留
                out.append('')
                out.append(f'# {text}')
                out.append('')
                continue
            if lvl and 2 <= lvl <= 4 and len(text) < 120:
                prefix = ''
                if num and num[0] and not text.startswith(num[0]):
                    prefix = num[0] + num[1]
                out.append('')
                out.append('#' * lvl + ' ' + prefix + text)
                out.append('')
                continue
            # 列表等带编号段落: 编号合成回文字前缀(无可见编号的空 lvlText 不补)
            if num and num[0] and not text.startswith(num[0]):
                text = num[0] + num[1] + text
            if not keep_hints:
                # 跨段提示区: 【开区(可在段中), 】收区
                if in_hint:
                    if '】' in text:
                        in_hint = False
                        tail = text.split('】', 1)[1].strip()
                        if tail:
                            out.append(tail)
                            out.append('')
                    continue
                # 单段完整提示 / 段中混合
                stripped = re.sub(r'【[^】]*】', '', text).strip()
                if not stripped:
                    continue
                if '【' in stripped and '】' not in stripped:
                    # 提示区从段中开始: 保留【前的前缀, 进入提示区
                    prefix = stripped.split('【', 1)[0].strip()
                    if prefix:
                        out.append(prefix)
                        out.append('')
                    in_hint = True
                    continue
                text = stripped
            out.append(text)
            out.append('')
        else:  # tbl
            tbl_idx += 1
            # 跳过封面表(含 发明名称 行)与注意事项表(含 注意事项)
            first_row = ''.join(t.text or '' for t in el.iter(qn('w:t')))[:30]
            if '发明名称' in first_row and tbl_idx == 0:
                continue
            if '注意事项' in first_row:
                continue
            md_rows = tbl_to_md(el)
            if md_rows:
                out.extend(md_rows)
                out.append('')
    md = '\n'.join(out)
    md = re.sub(r'\n{3,}', '\n\n', md).strip() + '\n'
    with open(out_md, 'w', encoding='utf-8') as f:
        f.write(md)
    return out_md, len(md)


def main():
    args = sys.argv[1:]
    if len(args) < 2:
        print('用法: docx2md.py <in.docx> <out.md> [--keep-hints]')
        sys.exit(1)
    keep = '--keep-hints' in args
    src, dst = args[0], args[1]
    out, n = convert(src, dst, keep)
    print(f'OK -> {out} ({n} chars)')


if __name__ == '__main__':
    main()
