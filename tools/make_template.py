# -*- coding: utf-8 -*-
"""make_template.py — 从用户定稿的交底书 docx 提取"空白格式模板"

用法:
    .venv_patent/Scripts/python.exe tools/make_template.py <参考docx> <输出模板docx>

作用: 把用户手工调好格式的交底书作为"格式基准", 保留其
  - 封面表(公司编号/发明名称/申报单位/申报类型/发明人/技术联系人) 与 注意事项表
  - 页首标题段("公司专利申请/技术交底书/〔集团名称〕")
  - 全部样式(styless.xml: Heading 1~4 黑体16/15/15/14pt、Body Text 宋体12pt 等)
  - 编号定义(numbering.xml: 一级中文编号、子节"N.%1"编号等)
  - 页脚(页码) 与 页面设置(A4/页边距)
清除: 正文全部内容(自"零、术语定义和解释"起的段落与表格); 封面"发明名称"值格清空(生成时回填);
      正文清空后残留的孤儿图片一并删除(保持模板与产物干净)。
"""
import sys
import docx
from docx.oxml.ns import qn

ZERO_MARKER = '零、术语定义和解释'


def _prune_unused_images(d):
    """删除正文未引用的图片关系(文档残留的孤儿图片, 保持模板干净)"""
    used = set()
    for blip in d.element.iter(qn('a:blip')):
        rid = blip.get(qn('r:embed')) or blip.get(qn('r:link'))
        if rid:
            used.add(rid)
    for imagedata in d.element.iter('{urn:schemas-microsoft-com:vml}imagedata'):
        rid = imagedata.get(qn('r:id'))
        if rid:
            used.add(rid)
    for rel in list(d.part.rels.values()):
        try:
            if rel.reltype.endswith('/image') and rel.rId not in used:
                d.part.drop_rel(rel.rId)
        except Exception:
            pass


def strip_to_template(src, dst):
    d = docx.Document(src)
    body = d.element.body
    found = False
    to_remove = []
    for child in list(body.iterchildren()):
        tag = child.tag.split('}')[1]
        if tag == 'sectPr':
            continue
        if tag == 'p':
            txt = ''.join(t.text or '' for t in child.iter(qn('w:t'))).strip()
            if not found and txt.startswith(ZERO_MARKER):
                found = True
            if found:
                to_remove.append(child)
        elif tag == 'tbl':
            if found:
                to_remove.append(child)
    for el in to_remove:
        body.remove(el)
    # 清空封面"发明名称"值格(生成时按案件回填)
    if d.tables:
        cover = d.tables[0]
        if len(cover.rows) > 1 and cover.rows[1].cells[0].text.strip() == '发明名称':
            cover.rows[1].cells[1].text = ''
    # 删除正文清空后残留的孤儿图片(保持模板与产物干净)
    _prune_unused_images(d)
    d.save(dst)
    print(f'OK -> {dst} | removed {len(to_remove)} block(s), zero_found={found}')


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print('用法: make_template.py <参考docx> <输出模板docx>')
        sys.exit(1)
    strip_to_template(sys.argv[1], sys.argv[2])
